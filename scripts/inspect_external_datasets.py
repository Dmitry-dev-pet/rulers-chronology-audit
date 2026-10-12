"""Inspect an external leadership dataset without importing historical claims.

The inspected file stays local; only structural metadata is printed. Canonical
Git data, SQLite, site output and historical conclusions are never modified.

CSV/TSV need only the standard library. Stata requires pandas, Excel requires
openpyxl. Those dependencies are optional and are never used during site build.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from pathlib import Path
from typing import Iterable

SOURCES = {
    "archigos": {
        "version": "4.1",
        "url": "https://www.rochester.edu/college/faculty/hgoemans/data.htm",
        "candidate_identifiers": ["obsid", "leadid"],
    },
    "plad": {
        "version": "verify_from_download",
        "url": "https://plad.uni-goettingen.de/data/",
        "candidate_identifiers": ["plad_id", "archigos_id", "idacr"],
    },
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def is_missing(val: object) -> bool:
    return val is None or (isinstance(val, float) and math.isnan(val)) or str(val).strip() in {"", "<NA>", "NaT"}


def inspect_rows(headers: list[str], rows: Iterable[dict], identifiers: list[str]) -> dict:
    if not headers or any(not isinstance(h, str) or not h.strip() for h in headers):
        raise ValueError("Missing or blank column names")
    if len(set(headers)) != len(headers):
        raise ValueError("Duplicate column names: refusing ambiguous data")
    present = [col for col in identifiers if col in headers]
    missing = {col: 0 for col in headers}
    seen = {col: set() for col in present}
    duplicate = {col: 0 for col in present}
    count = 0
    for row in rows:
        count += 1
        for col in headers:
            val = row.get(col)
            if is_missing(val):
                missing[col] += 1
        for col in present:
            val = row.get(col)
            if is_missing(val):
                continue
            key = str(val)
            if key in seen[col]:
                duplicate[col] += 1
            else:
                seen[col].add(key)
    if count == 0:
        raise ValueError("Dataset has no data rows")
    return {
        "rows": count,
        "columns": headers,
        "missing_values_by_column": missing,
        "candidate_identifiers_present": present,
        "repeated_identifier_values": duplicate,
    }


def csv_table(path: Path, delimiter: str) -> tuple[list[str], Iterable[dict]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream, delimiter=delimiter)
        headers = list(reader.fieldnames or [])
        rows = list(reader)
    if any(None in row for row in rows):
        raise ValueError("CSV row has more values than its header")
    return headers, rows


def stata_table(path: Path) -> tuple[list[str], Iterable[dict]]:
    try:
        import pandas as pd
    except ImportError as exc:
        raise RuntimeError("Stata .dta inspection requires 'pip install pandas'") from exc
    frame = pd.read_stata(path, convert_categoricals=False, convert_dates=False)
    headers = list(frame.columns)
    return headers, frame.to_dict(orient="records")


def excel_table(path: Path, sheet: str | None) -> tuple[list[str], Iterable[dict]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel .xlsx inspection requires 'pip install openpyxl'") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet and sheet not in workbook.sheetnames:
        workbook.close()
        raise ValueError("Unknown Excel sheet: " + sheet)
    try:
        worksheet = workbook[sheet] if sheet else workbook.active
        values = worksheet.iter_rows(values_only=True)
        headers = [str(v).strip() if v is not None else "" for v in next(values, ())]
        rows = []
        for values_row in values:
            if len(values_row) > len(headers) and any(
                v is not None for v in values_row[len(headers):]
            ):
                raise ValueError("Excel row exceeds header width")
            rows.append(dict(zip(headers, values_row)))
    finally:
        workbook.close()
    return headers, rows


def inspect_file(dataset: str, path: Path, *, sheet: str | None = None) -> dict:
    if dataset not in SOURCES:
        raise ValueError("Unknown source: " + dataset)
    if not path.is_file():
        raise FileNotFoundError(path)
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        headers, rows = csv_table(path, "\t" if suffix == ".tsv" else ",")
    elif suffix == ".dta":
        headers, rows = stata_table(path)
    elif suffix == ".xlsx":
        headers, rows = excel_table(path, sheet)
    else:
        raise ValueError("Unsupported file type. Extract archives before inspection: " + suffix)
    structure = inspect_rows(headers, rows, SOURCES[dataset]["candidate_identifiers"])
    return {
        "dataset": dataset,
        "version": SOURCES[dataset]["version"],
        "source_url": SOURCES[dataset]["url"],
        "file_name": path.name,
        "file_bytes": path.stat().st_size,
        "sha256": sha256_file(path),
        "status": "structure_inspected_only",
        "canonical_data_modified": False,
        "historical_truth_verified_by_ci": False,
        **structure,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("dataset", choices=sorted(SOURCES))
    parser.add_argument("file", type=Path)
    parser.add_argument("--sheet", help="Excel sheet name; default first worksheet")
    args = parser.parse_args()
    print(json.dumps(inspect_file(args.dataset, args.file, sheet=args.sheet),
                     ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
